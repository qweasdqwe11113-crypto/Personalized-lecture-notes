"""Exercise parsing, persistence, download and web-import boundaries in a disposable library."""
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from flask import Flask
from uploads import ImportError, WebText, extract, public_target, register_uploads


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        app = Flask(__name__)
        app.testing = True
        register_uploads(app, self.directory.name)
        self.client = app.test_client()

    def tearDown(self):
        self.directory.cleanup()

    def test_file_roundtrip_and_download(self):
        text = '# 教学资料\n等式两边同加同减。'
        response = self.client.post('/api/uploads', data={'file': (BytesIO(text.encode()), '资料.md')})
        self.assertEqual(response.status_code, 201)
        item = response.json['upload']
        self.assertEqual(item['title'], '资料')
        self.assertEqual(self.client.get('/api/uploads/' + item['id']).json['upload']['content'], text)
        self.assertEqual(self.client.get('/api/uploads/' + item['id'] + '/download').data, text.encode())
        self.assertEqual(len(self.client.get('/api/uploads').json['uploads']), 1)
        self.assertNotIn('raw', self.client.get('/api/uploads').json['uploads'][0])

    def test_invalid_files_do_not_persist(self):
        for raw, name in [(b'', 'empty.txt'), (b'<script>x</script>', 'bad.exe'), (b'broken', 'bad.docx')]:
            response = self.client.post('/api/uploads', data={'file': (BytesIO(raw), name)})
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.get('/api/uploads').json['uploads'], [])

    def test_docx_paragraph_and_table(self):
        from docx import Document
        doc = Document(); doc.add_paragraph('知识点讲解')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格中的内容'
        buffer = BytesIO(); doc.save(buffer)
        content = extract(buffer.getvalue(), 'test.docx')
        self.assertIn('知识点讲解', content)
        self.assertIn('表格中的内容', content)

    def test_url_import_extracts_text_without_script(self):
        body = '<title>课程文章</title><p>等式性质</p><script>danger()</script><p>课堂例题</p>'.encode()
        with patch('uploads.fetch_public', return_value=(body, 'text/html', 'utf-8', 'https://example.com/lesson')):
            response = self.client.post('/api/uploads', json={'url': 'https://example.com/lesson'})
        self.assertEqual(response.status_code, 201)
        item = response.json['upload']
        self.assertEqual(item['title'], '课程文章')
        content = self.client.get('/api/uploads/' + item['id']).json['upload']['content']
        self.assertIn('等式性质', content)
        self.assertNotIn('danger', content)

    def test_nonpublic_or_credential_url_rejected(self):
        for url in ['file:///secret', 'https://user:pass@example.com', 'http://example.com:9000']:
            with self.assertRaises(ImportError): public_target(url)
        with patch('uploads.socket.getaddrinfo', return_value=[(2, 1, 6, '', ('127.0.0.1', 80))]):
            with self.assertRaises(ImportError): public_target('http://localhost')
        self.assertEqual(self.client.post('/api/uploads', json=[]).status_code, 400)
        self.assertEqual(self.client.get('/api/uploads/missing').status_code, 404)


if __name__ == '__main__':
    unittest.main()
