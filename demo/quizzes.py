"""Persistent, bounded question generation with frozen requests and independent exports."""
from collections import deque
from hashlib import sha256
import html
from io import BytesIO
import json
from pathlib import Path
from threading import Condition, Thread
from urllib.parse import quote
from uuid import uuid4

from flask import Response, jsonify, request
from psycopg.types.json import Jsonb

from prompts import PreparedRequest
from qwen import ModelError

TYPES = {'single_choice': '单项选择', 'fill_blank': '填空', 'short_answer': '简答',
         'calculation': '计算', 'application': '综合应用'}
DIFFICULTIES = {'basic': '基础', 'medium': '中等', 'advanced': '拔高', 'mixed': '混合难度'}
LEVELS = {'profile': '按学生画像', 'foundation': '基础待巩固', 'standard': '常规水平', 'advanced': '学有余力'}


class QuizError(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def validate_options(data, knowledge_ids):
    if not isinstance(data, dict):
        raise QuizError('请提交有效的出题设置。')
    count = data.get('count', 10)
    if type(count) is not int or not 1 <= count <= 30:
        raise QuizError('出题数量请填写 1—30 的整数。')
    difficulty, level = data.get('difficulty', 'medium'), data.get('student_level', 'standard')
    if not isinstance(difficulty, str) or difficulty not in DIFFICULTIES:
        raise QuizError('请选择基础、中等、拔高或混合难度。')
    if not isinstance(level, str) or level not in LEVELS:
        raise QuizError('请选择有效的学生水平。')
    types = data.get('types', ['single_choice', 'short_answer'])
    if (not isinstance(types, list) or not types or len(types) > len(TYPES)
            or any(not isinstance(t, str) or t not in TYPES for t in types) or len(set(types)) != len(types)):
        raise QuizError('请至少选择一种有效题型。')
    focus = data.get('knowledge_ids', [])
    if (not isinstance(focus, list) or any(type(k) is not int or k not in knowledge_ids for k in focus)
            or len(set(focus)) != len(focus)):
        raise QuizError('请选择当前副本内的知识点。')
    note = data.get('requirements', '')
    if not isinstance(note, str) or len(note.strip()) > 2000:
        raise QuizError('出题要求请控制在 2000 字以内。')
    return dict(count=count, difficulty=difficulty, student_level=level, types=types,
                knowledge_ids=focus or list(knowledge_ids), requirements=note.strip())


def question_plan(options, states):
    count = options['count']
    if options['difficulty'] == 'mixed':
        # Largest remainder allocation keeps 50/30/20 proportions and the exact requested total.
        allocation = [count * n // 10 for n in (5, 3, 2)]
        order = sorted(range(3), key=lambda i: -(count * (5, 3, 2)[i] % 10))
        for i in order[:count - sum(allocation)]:
            allocation[i] += 1
        difficulties = [d for d, n in zip(('basic', 'medium', 'advanced'), allocation) for _ in range(n)]
    else:
        difficulties = [options['difficulty']] * count
    focus = list(options['knowledge_ids'])
    if options['student_level'] == 'profile':
        rank = {s['id']: {'needs_support': 0, 'partial': 1, 'unknown': 2, 'mastered': 3}.get(s['state'], 2) for s in states}
        focus.sort(key=lambda k: rank.get(k, 2))
    return [dict(number=i + 1, type=options['types'][i % len(options['types'])],
                 difficulty=difficulties[i], knowledge_id=focus[i % len(focus)]) for i in range(count)]


def prepare_quiz(context, options, model):
    payload = dict(model=model, temperature=0.4, max_tokens=16000,
                   chat_template_kwargs={'enable_thinking': False}, messages=[
        {'role': 'system', 'content': (Path(__file__).with_name('prompts') / 'quiz-v1.md').read_text(encoding='utf-8')},
        {'role': 'user', 'content': json.dumps(dict(context=context, options=options), ensure_ascii=False, default=str)}])
    body = json.dumps(payload, ensure_ascii=False, separators=(',', ':'))
    return PreparedRequest(body, sha256(body.encode()).hexdigest(), 'quiz-v1', 'quiz-context-v1')


def parse_questions(content, context):
    def invalid(message):
        raise ModelError('invalid_questions', message, 'output')
    try:
        result = json.loads(content)
    except (json.JSONDecodeError, TypeError):
        invalid('模型未返回有效题目 JSON，请重试。')
    plan = context['question_plan']
    if not isinstance(result, dict) or set(result) != {'questions'} or not isinstance(result['questions'], list) or len(result['questions']) != len(plan):
        invalid('模型返回题数与出题计划不符，未保存为完整题组。')
    sources = {f'M{m["id"]}' for m in context['materials']}
    seen = set()
    fields = {'number', 'type', 'difficulty', 'knowledge_id', 'stem', 'options', 'answer', 'analysis', 'common_mistake', 'source_ids'}
    for q, expected in zip(result['questions'], plan):
        if not isinstance(q, dict) or set(q) != fields:
            invalid('题目字段不完整，未保存为完整题组。')
        if any(type(q[k]) is not type(v) or q[k] != v for k, v in expected.items()):
            invalid('模型未遵守约定的题型、难度或知识点分配，请重试。')
        for field, maximum in [('stem', 3000), ('answer', 2500), ('analysis', 5000), ('common_mistake', 1500)]:
            if not isinstance(q[field], str) or not q[field].strip() or len(q[field]) > maximum:
                invalid('题干、答案或解析不完整或过长，请重试。')
        stem = ''.join(q['stem'].split()).casefold()
        if stem in seen:
            invalid('题组存在重复题干，请重试。')
        seen.add(stem)
        choices = q['options']
        if q['type'] == 'single_choice':
            if (not isinstance(choices, list) or len(choices) != 4
                    or any(not isinstance(v, str) or not v.strip() or len(v) > 1500 for v in choices)
                    or len({v.strip().casefold() for v in choices}) != 4 or q['answer'] not in ('A', 'B', 'C', 'D')):
                invalid('单选题的四个选项或答案格式无效，请重试。')
        elif choices != []:
            invalid('非选择题不应包含选择选项，请重试。')
        ids = q['source_ids']
        if (not isinstance(ids, list) or any(not isinstance(v, str) or v not in sources for v in ids)
                or len(set(ids)) != len(ids)):
            invalid('题目引用了不存在的材料编号，请重试。')
    return result


class QuizQueue:
    def __init__(self, db, call_model, generation_lock, capacity=4):
        self.db, self.call_model, self.lock, self.capacity = db, call_model, generation_lock, capacity
        self.condition, self.pending = Condition(), deque()
        self.active = None
        self.thread = None
        self.stopping = False

    def recover(self):
        with self.db() as conn:
            conn.execute("UPDATE quiz_jobs SET status='interrupted',finished_at=now(),error_message='本地服务已重启，可按原设置重试。' WHERE status IN ('queued','running')")

    def detail(self, job_id, full=True):
        with self.db() as conn:
            row = conn.execute('SELECT * FROM quiz_jobs WHERE id=%s', (job_id,)).fetchone()
        if not row:
            raise QuizError('出题记录不存在。', 404)
        row['integrity_valid'] = sha256(row['request_body'].encode()).hexdigest() == row['request_sha256']
        if not full:
            row.pop('request_body')
        with self.condition:
            row['confirmation_pending'] = row['status'] in ('queued', 'running') and row['id'] != self.active and not any(p[0] == job_id for p in self.pending)
        return row

    def submit(self, topic, sid, context, options, prepared, parent=None):
        with self.condition:
            if self.stopping:
                raise QuizError('出题服务正在停止，请稍后重试。', 503)
            if len(self.pending) + int(self.active is not None) >= self.capacity:
                raise QuizError('出题队列已满，请等待当前任务完成。', 429)
            job_id = uuid4().hex
            with self.db() as conn:
                conn.execute('''INSERT INTO quiz_jobs(id,parent_id,topic_key,student_id,context,options,request_body,request_sha256)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s)''',
                    (job_id, parent, topic, sid, Jsonb(context), Jsonb(options), prepared.body, prepared.sha256))
            self.pending.append((job_id, context, prepared))
            if self.thread is None:
                self.thread = Thread(target=self._work, daemon=True, name='quiz-worker')
                self.thread.start()
            self.condition.notify_all()
        return self.detail(job_id, False)

    def retry(self, job_id):
        row = self.detail(job_id)
        if row['status'] not in ('failed', 'interrupted'):
            raise QuizError('仅失败或中断的题组可以按原输入重试。', 409)
        if not row['integrity_valid']:
            raise QuizError('原输入校验不一致，不能重试。', 409)
        prepared = PreparedRequest(row['request_body'], row['request_sha256'], 'quiz-v1', 'quiz-context-v1')
        return self.submit(row['topic_key'], row['student_id'], row['context'], row['options'], prepared, job_id)

    def _work(self):
        while True:
            with self.condition:
                self.condition.wait_for(lambda: self.pending or self.stopping)
                if self.stopping and not self.pending:
                    return
                job_id, context, prepared = self.pending.popleft()
                self.active = job_id
            try:
                with self.lock:
                    with self.db() as conn:
                        conn.execute("UPDATE quiz_jobs SET status='running',started_at=now() WHERE id=%s", (job_id,))
                    output = self.call_model(prepared, lambda: None)
                    result = parse_questions(output.content, context)
                    with self.db() as conn:
                        conn.execute("UPDATE quiz_jobs SET status='succeeded',result=%s,usage=%s,finished_at=now() WHERE id=%s",
                                     (Jsonb(result), Jsonb(output.usage), job_id))
            except Exception as exc:
                message = str(exc) if isinstance(exc, ModelError) else '出题未完成，请检查本地服务后重试。'
                try:
                    with self.db() as conn:
                        conn.execute("UPDATE quiz_jobs SET status='failed',error_message=%s,finished_at=now() WHERE id=%s", (message, job_id))
                except Exception:
                    pass
            finally:
                with self.condition:
                    self.active = None
                    self.condition.notify_all()

    def wait_idle(self, timeout=5):
        with self.condition:
            return self.condition.wait_for(lambda: self.active is None and not self.pending, timeout)

    def close(self):
        with self.condition:
            self.stopping = True
            self.condition.notify_all()
        if self.thread:
            self.thread.join(timeout=5)


def quiz_markdown(row, teacher=False):
    context = row['context']
    lines = ['# '+context['topic']['title']+' · 练习题', '',
             f"共 {row['options']['count']} 题 · {DIFFICULTIES[row['options']['difficulty']]}", '',
             '教师版 · 答案与解析需教师复核。' if teacher else '姓名：____________    日期：____________', '']
    for q in row['result']['questions']:
        lines.extend([f"## 第 {q['number']} 题 · {TYPES[q['type']]} · {DIFFICULTIES[q['difficulty']]}", '', q['stem'], ''])
        lines.extend(f'{letter}. {value}' for letter, value in zip('ABCD', q['options']))
        lines.extend(['', '作答：', '____________________________', ''])
    if teacher:
        lines.extend(['## 参考答案与解析', ''])
        names = {s['id']: s['name'] for s in context['states']}
        names.update(context.get('knowledge_names', {}))
        for q in row['result']['questions']:
            lines.extend([f"### 第 {q['number']} 题", f"知识点：{names.get(str(q['knowledge_id']), names.get(q['knowledge_id'], ''))}",
                          '参考答案：'+q['answer'], '', '解析：'+q['analysis'], '', '易错点：'+q['common_mistake'],
                          '材料依据：'+('、'.join(q['source_ids']) if q['source_ids'] else '自编题目，未标注直接材料依据，请复核。'), ''])
    return '\n'.join(lines)


def register_quizzes(app, db, queue, catalog, context_builder, model):
    @app.errorhandler(QuizError)
    def quiz_error(exc):
        return jsonify(error=str(exc)), exc.status

    @app.post('/api/quizzes')
    def create_quiz():
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            raise QuizError('请提交有效的出题设置。')
        key, sid = data.get('topic'), data.get('student_id')
        topics = catalog()
        if not isinstance(key, str) or key not in topics:
            raise QuizError('请选择有效的备课副本。')
        topic = topics[key]
        options = validate_options(data, topic['knowledge_ids'])
        if sid is not None and (type(sid) is not int or sid not in topic['student_ids']):
            raise QuizError('请选择可用于当前副本的学生。')
        if options['student_level'] == 'profile' and sid is None:
            raise QuizError('按学生画像出题时，请先选择学生。')
        with db() as conn:
            points = conn.execute('SELECT id,name FROM knowledge_points WHERE id=ANY(%s)', (topic['knowledge_ids'],)).fetchall()
            if options['student_level'] != 'profile':
                materials = conn.execute('SELECT * FROM materials WHERE knowledge_id=ANY(%s) ORDER BY id', (topic['knowledge_ids'],)).fetchall()
        if options['student_level'] == 'profile':
            ctx = context_builder(sid, key, 'profile')
            if not ctx:
                raise QuizError('无法读取学生档案，请刷新后重试。')
        else:
            ctx = dict(student={'name': '通用练习', 'profile': '未提供学生画像。'}, topic=topic,
                       states=[], materials=materials, personalization_mode='none')
        ctx['materials'] = [{**m, 'citation_id': f'M{m["id"]}'} for m in ctx['materials']]
        ctx['knowledge_names'] = {str(p['id']): p['name'] for p in points}
        ctx['question_plan'] = question_plan(options, ctx['states'])
        return jsonify(queue().submit(key, sid, ctx, options, prepare_quiz(ctx, options, model()))), 202

    @app.get('/api/quizzes')
    def list_quizzes():
        with db() as conn:
            rows = conn.execute('''SELECT id,status,created_at,options,error_message,context->'topic'->>'title' AS title,
                context->'student'->>'name' AS student_name FROM quiz_jobs ORDER BY created_at DESC LIMIT 50''').fetchall()
        return jsonify(quizzes=rows)

    @app.get('/api/quizzes/<job_id>')
    def get_quiz(job_id):
        return jsonify(queue().detail(job_id, False))

    @app.post('/api/quizzes/<job_id>/retry')
    def retry_quiz(job_id):
        return jsonify(queue().retry(job_id)), 202

    @app.get('/api/quizzes/<job_id>/download')
    def download_quiz(job_id):
        row = queue().detail(job_id, False)
        if row['status'] != 'succeeded' or not row['integrity_valid']:
            raise QuizError('题组尚未完成或输入校验失败，暂不能下载。', 409)
        version, fmt = request.args.get('version', 'student'), request.args.get('format', 'docx')
        if version not in ('student', 'teacher') or fmt not in ('md', 'txt', 'html', 'docx'):
            raise QuizError('请选择有效的导出版本与格式。')
        text = quiz_markdown(row, version == 'teacher')
        if fmt == 'html':
            content = '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>练习题</title><style>body{max-width:850px;margin:32px auto;padding:0 24px;font:16px/1.9 sans-serif}pre{white-space:pre-wrap;font:inherit}</style><body><pre>'+html.escape(text)+'</pre></body></html>'
            mime = 'text/html; charset=utf-8'
        elif fmt == 'docx':
            from docx import Document
            document = Document()
            for line in text.splitlines():
                if line.startswith('#'):
                    level = len(line) - len(line.lstrip('#'))
                    document.add_heading(line[level:].strip(), min(level, 3))
                else:
                    document.add_paragraph(line)
            stream = BytesIO()
            document.save(stream)
            content, mime = stream.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            content, mime = text, 'text/plain; charset=utf-8' if fmt == 'txt' else 'text/markdown; charset=utf-8'
        filename = f"练习题-{'教师版' if version == 'teacher' else '学生版'}.{fmt}"
        return Response(content, content_type=mime, headers={'Content-Disposition': f"attachment; filename=quiz-{job_id}.{fmt}; filename*=UTF-8''{quote(filename)}"})
