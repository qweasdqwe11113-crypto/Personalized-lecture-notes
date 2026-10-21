"""Local document library and bounded public-web imports, independent of model calls."""
from contextlib import closing
from datetime import datetime, timezone
from html.parser import HTMLParser
from io import BytesIO
import http.client
import ipaddress
from pathlib import Path
import socket
import sqlite3
import ssl
import uuid
from urllib.parse import urljoin, urlsplit, unquote, quote
from zipfile import ZipFile, BadZipFile

from flask import jsonify, request, send_file

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 200000
EXTENSIONS = {'.txt', '.md', '.pdf', '.docx'}


class ImportError(ValueError):
    pass


class WebText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.title = [], []
        self.ignored = 0
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style', 'noscript', 'svg'}:
            self.ignored += 1
        if tag == 'title':
            self.in_title = True
        if tag in {'p', 'br', 'div', 'li', 'h1', 'h2', 'h3', 'article', 'section'}:
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in {'script', 'style', 'noscript', 'svg'}:
            self.ignored = max(0, self.ignored - 1)
        if tag == 'title':
            self.in_title = False
        if tag in {'p', 'div', 'li', 'h1', 'h2', 'h3'}:
            self.parts.append('\n')

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
        elif not self.ignored:
            self.parts.append(data)


def decode_text(raw, charset=None):
    for encoding in dict.fromkeys([charset, 'utf-8-sig', 'gb18030']):
        if not encoding:
            continue
        try:
            return raw.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            pass
    raise ImportError('无法识别文字编码，请将文件保存为 UTF-8 后重试。')


def extract(raw, filename):
    ext = Path(filename).suffix.lower()
    if ext not in EXTENSIONS:
        raise ImportError('支持 TXT、Markdown、PDF 和 DOCX 文件。')
    if ext in {'.txt', '.md'}:
        return decode_text(raw)
    try:
        if ext == '.docx':
            from docx import Document
            with ZipFile(BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                    raise ImportError('文档解压后过大，请拆分文件。')
            doc = Document(BytesIO(raw))
            return '\n'.join([p.text for p in doc.paragraphs] + [
                ' | '.join(cell.text for cell in row.cells)
                for table in doc.tables for row in table.rows])
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise ImportError('请先解密 PDF，再上传。')
        if len(reader.pages) > 200:
            raise ImportError('PDF 最多支持 200 页，请拆分后上传。')
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    except ImportError:
        raise
    except Exception as exc:
        raise ImportError('文件解析失败，请确认文件完整且格式正确。') from exc


def public_target(url):
    try:
        parsed = urlsplit(url)
        if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
            raise ImportError('请输入公开网页的 HTTP 或 HTTPS 网址。')
        if parsed.port not in {None, 80, 443}:
            raise ImportError('仅支持标准 HTTP / HTTPS 端口。')
        host = parsed.hostname.encode('idna').decode('ascii')
        port = parsed.port or (443 if parsed.scheme == 'https' else 80)
        addresses = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
            raise ImportError('仅支持公开互联网网址，不能导入本机或内网地址。')
        return parsed, host, port, addresses[0][4][0]
    except ImportError:
        raise
    except (ValueError, UnicodeError, socket.gaierror) as exc:
        raise ImportError('网址无效或无法解析，请检查网址。') from exc


def fetch_public(url):
    # Connect to the validated IP; retain the original host for TLS and Host.
    # Revalidate each redirect and never forward cookies or credentials.
    for _ in range(4):
        parsed, host, port, address = public_target(url)
        conn = http.client.HTTPConnection(host, port, timeout=15)
        try:
            sock = socket.create_connection((address, port), timeout=15)
            if parsed.scheme == 'https':
                try:
                    sock = ssl.create_default_context().wrap_socket(sock, server_hostname=host)
                except Exception:
                    sock.close()
                    raise
            conn.sock = sock
            target = quote(parsed.path or '/', safe="/%:@!$&()*+,;=-._~")
            if parsed.query:
                target += '?' + quote(parsed.query, safe="=&%/:?@!$()*+,;~-._")
            conn.request('GET', target, headers={'User-Agent': 'Zhiye-Material-Importer/1.0', 'Accept-Encoding': 'identity'})
            response = conn.getresponse()
            if response.status in {301, 302, 303, 307, 308}:
                location = response.getheader('Location')
                if not location:
                    raise ImportError('网页重定向缺少目标网址。')
                url = urljoin(url, location)
                continue
            if response.status != 200:
                raise ImportError(f'网页返回 HTTP {response.status}，请使用可直接访问的网址。')
            raw = response.read(MAX_BYTES + 1)
            if len(raw) > MAX_BYTES:
                raise ImportError('网页或文件超过 10 MB，请使用更小的资料。')
            mime = response.headers.get_content_type()
            charset = response.headers.get_content_charset()
            return raw, mime, charset, url
        except (OSError, http.client.HTTPException) as exc:
            raise ImportError('网址读取失败，请检查网络或换一个可公开访问的网址。') from exc
        finally:
            conn.close()
    raise ImportError('网页重定向次数过多，请使用最终网址。')


def read_upload(root, upload_id):
    """Read one stored upload so retrieval can index it; None when it is not there."""
    store = Path(root) / '.runtime' / 'uploads.sqlite3'
    if not store.exists():
        return None
    with closing(sqlite3.connect(store, timeout=10)) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute('SELECT id,title,filename,source_url,content FROM uploads WHERE id = ?',
                           (upload_id,)).fetchone()
    return dict(row) if row else None


def register_uploads(app, root):
    store = Path(root) / '.runtime' / 'uploads.sqlite3'

    def connect():
        store.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(store, timeout=10)
        conn.row_factory = sqlite3.Row
        conn.execute('''CREATE TABLE IF NOT EXISTS uploads (
            id TEXT PRIMARY KEY, title TEXT NOT NULL, filename TEXT NOT NULL,
            kind TEXT NOT NULL, source_url TEXT NOT NULL, content TEXT NOT NULL,
            raw BLOB NOT NULL, created_at TEXT NOT NULL)''')
        conn.commit()
        return conn

    def summary(row):
        return {key: row[key] for key in ('id', 'title', 'filename', 'kind', 'source_url', 'created_at')}

    @app.get('/api/uploads')
    def list_uploads():
        with closing(connect()) as conn:
            rows = conn.execute('SELECT id,title,filename,kind,source_url,created_at,length(content) AS characters FROM uploads ORDER BY created_at DESC LIMIT 100').fetchall()
        return jsonify(uploads=[{**summary(row), 'characters': row['characters']} for row in rows])

    @app.post('/api/uploads')
    def create_upload():
        try:
            if request.mimetype == 'multipart/form-data':
                file = request.files.get('file')
                if not file or not file.filename:
                    raise ImportError('请选择要上传的资料文件。')
                filename = file.filename.replace('\\', '/').rsplit('/', 1)[-1][:180]
                raw = file.read(MAX_BYTES + 1)
                if not raw or len(raw) > MAX_BYTES:
                    raise ImportError('文件不能为空，且不能超过 10 MB。')
                content = extract(raw, filename)
                title = request.form.get('title', '').strip() or Path(filename).stem
                kind, source_url = 'file', ''
            else:
                data = request.get_json(silent=True) or {}
                if not isinstance(data, dict) or not isinstance(data.get('url'), str):
                    raise ImportError('请输入要导入的网址。')
                url = data['url'].strip()
                if len(url) > 2048:
                    raise ImportError('网址过长。')
                raw, mime, charset, source_url = fetch_public(url)
                filename = unquote(urlsplit(source_url).path.rsplit('/', 1)[-1])[:180] or 'webpage.html'
                if mime in {'text/html', 'application/xhtml+xml'}:
                    parser = WebText()
                    parser.feed(decode_text(raw, charset))
                    content = '\n'.join(line.strip() for line in ''.join(parser.parts).splitlines() if line.strip())
                    title = ''.join(parser.title).strip() or urlsplit(source_url).hostname
                    filename = 'webpage.html'
                elif mime.startswith('text/'):
                    content, title = decode_text(raw, charset), filename
                elif Path(filename).suffix.lower() in {'.pdf', '.docx'}:
                    content, title = extract(raw, filename), Path(filename).stem
                else:
                    raise ImportError('网址内容不支持，请使用网页、文本、PDF 或 DOCX 的直接链接。')
                custom_title = data.get('title', '')
                if not isinstance(custom_title, str):
                    raise ImportError('资料名称格式不正确。')
                title = custom_title.strip() or title
                kind = 'url'
            content = content.strip()
            if not content or '\x00' in content:
                raise ImportError('未提取到有效文字。扫描版 PDF 请先进行文字识别。')
            if len(content) > MAX_TEXT:
                raise ImportError('资料正文超过 20 万字符，请拆分后上传。')
            if len(title) > 120:
                raise ImportError('资料名称最多 120 字。')
            record = dict(id=uuid.uuid4().hex, title=title, filename=filename, kind=kind,
                          source_url=source_url, content=content, raw=raw,
                          created_at=datetime.now(timezone.utc).isoformat())
            with closing(connect()) as conn, conn:
                conn.execute('INSERT INTO uploads VALUES (:id,:title,:filename,:kind,:source_url,:content,:raw,:created_at)', record)
            return jsonify(upload={**summary(record), 'characters': len(content)}), 201
        except ImportError as exc:
            return jsonify(error=str(exc)), 400

    @app.get('/api/uploads/<upload_id>')
    def upload_detail(upload_id):
        with closing(connect()) as conn:
            row = conn.execute('SELECT * FROM uploads WHERE id=?', (upload_id,)).fetchone()
        if row is None:
            return jsonify(error='资料不存在。'), 404
        return jsonify(upload={**summary(row), 'content': row['content'], 'characters': len(row['content'])})

    @app.get('/api/uploads/<upload_id>/download')
    def download_upload(upload_id):
        with closing(connect()) as conn:
            row = conn.execute('SELECT filename,raw FROM uploads WHERE id=?', (upload_id,)).fetchone()
        if row is None:
            return jsonify(error='资料不存在。'), 404
        return send_file(BytesIO(row['raw']), as_attachment=True, download_name=row['filename'], mimetype='application/octet-stream')
